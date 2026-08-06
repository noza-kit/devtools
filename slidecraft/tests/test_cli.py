from pathlib import Path

from pptx import Presentation
from typer.testing import CliRunner

from slidecraft.cli import app

runner = CliRunner()


def test_get_prompt_flag_prints_prompt_and_exits():
    result = runner.invoke(app, ["--get-prompt"])
    assert result.exit_code == 0
    assert 'class="presentation"' in result.output


def test_no_args_prints_help():
    result = runner.invoke(app, [])
    assert result.exit_code == 0
    assert "generate" in result.output.lower() or "Usage" in result.output


def test_generate_creates_pptx(sample_html: Path, tmp_path: Path):
    out = tmp_path / "output.pptx"
    result = runner.invoke(app, ["generate", str(sample_html), "-o", str(out)])
    assert result.exit_code == 0, result.output
    assert out.exists()
    prs = Presentation(out)
    assert len(prs.slides) == 2


def test_generate_with_template(sample_html: Path, tmp_path: Path):
    template = tmp_path / "template.pptx"
    Presentation().save(template)
    out = tmp_path / "output.pptx"
    result = runner.invoke(app, ["generate", str(sample_html), "--template", str(template), "-o", str(out)])
    assert result.exit_code == 0, result.output
    assert out.exists()


def test_generate_missing_html_fails(tmp_path: Path):
    result = runner.invoke(app, ["generate", str(tmp_path / "missing.html"), "-o", str(tmp_path / "out.pptx")])
    assert result.exit_code == 1
    assert "エラー" in result.output


def test_generate_missing_template_fails(sample_html: Path, tmp_path: Path):
    result = runner.invoke(
        app, ["generate", str(sample_html), "--template", str(tmp_path / "missing.pptx"), "-o", str(tmp_path / "out.pptx")]
    )
    assert result.exit_code == 1
    assert "エラー" in result.output


def test_generate_invalid_html_reports_parse_error(tmp_path: Path):
    html_path = tmp_path / "bad.html"
    html_path.write_text("<div>no slide sections here</div>", encoding="utf-8")
    result = runner.invoke(app, ["generate", str(html_path), "-o", str(tmp_path / "out.pptx")])
    assert result.exit_code == 1
    assert "エラー" in result.output
