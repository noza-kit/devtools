from pathlib import Path

from pptx import Presentation

from slidecraft.builder import build_presentation
from slidecraft.parser import ContentItem, ImageRef, SlideData, TextRun, parse_html


def test_build_creates_one_slide_per_section(sample_html: Path):
    slides = parse_html(sample_html)
    prs = build_presentation(slides)
    assert len(prs.slides) == 2


def test_title_set_on_slide(sample_html: Path):
    slides = parse_html(sample_html)
    prs = build_presentation(slides)
    assert prs.slides[0].shapes.title.text == "市場調査のサマリー"


def test_bold_run_preserved(sample_html: Path):
    slides = parse_html(sample_html)
    prs = build_presentation(slides)
    body = [p for p in prs.slides[0].placeholders if p.placeholder_format.idx != 0][0]
    first_run = body.text_frame.paragraphs[0].runs[0]
    assert first_run.text == "ターゲット層:"
    assert first_run.font.bold is True


def test_nested_bullet_level_preserved(sample_html: Path):
    slides = parse_html(sample_html)
    prs = build_presentation(slides)
    body = [p for p in prs.slides[0].placeholders if p.placeholder_format.idx != 0][0]
    assert body.text_frame.paragraphs[2].level == 1


def test_image_slide_has_picture_shape(sample_html: Path):
    slides = parse_html(sample_html)
    prs = build_presentation(slides)
    shape_types = [type(s).__name__ for s in prs.slides[0].shapes]
    assert "Picture" in shape_types


def test_text_only_slide_has_no_picture(sample_html: Path):
    slides = parse_html(sample_html)
    prs = build_presentation(slides)
    shape_types = [type(s).__name__ for s in prs.slides[1].shapes]
    assert "Picture" not in shape_types


def test_text_placeholder_survives_alongside_image(sample_html: Path):
    """Regression test: an earlier bug compared placeholder objects by
    `is` identity to decide which one to keep for text vs. which to
    remove for the image, but `slide.placeholders` hands back a fresh
    wrapper object on every access, so the identity check silently never
    matched -- the placeholder that had just been filled with text was
    deleted, not the other (empty) one."""
    slides = parse_html(sample_html)
    prs = build_presentation(slides)
    body = [p for p in prs.slides[0].placeholders if p.placeholder_format.idx != 0][0]
    assert body.text_frame.paragraphs[0].runs[0].text == "ターゲット層:"


def test_notes_set(sample_html: Path):
    slides = parse_html(sample_html)
    prs = build_presentation(slides)
    assert "重要なスライドです" in prs.slides[0].notes_slide.notes_text_frame.text


def test_slide_without_notes_has_no_notes_text(sample_html: Path):
    slides = parse_html(sample_html)
    prs = build_presentation(slides)
    # notes_slide is lazily created by python-pptx; only assert if present
    if prs.slides[1].has_notes_slide:
        assert prs.slides[1].notes_slide.notes_text_frame.text == ""


def test_build_with_custom_template(tmp_path: Path, sample_html: Path):
    custom_template = tmp_path / "template.pptx"
    Presentation().save(custom_template)  # a minimal valid .pptx to use as a template
    slides = parse_html(sample_html)
    prs = build_presentation(slides, template_path=custom_template)
    assert len(prs.slides) == 2


def test_multiple_images_stack_vertically(tmp_path: Path, sample_image: Path):
    data = SlideData(
        title="Two images",
        content_items=[],
        images=[ImageRef(path=sample_image), ImageRef(path=sample_image)],
    )
    prs = build_presentation([data])
    pictures = [s for s in prs.slides[0].shapes if type(s).__name__ == "Picture"]
    assert len(pictures) == 2
    assert pictures[0].top != pictures[1].top


def test_slide_with_no_content_items_and_no_images(tmp_path: Path):
    data = SlideData(title="Empty body", content_items=[], images=[])
    prs = build_presentation([data])
    assert prs.slides[0].shapes.title.text == "Empty body"


def test_image_preserves_aspect_ratio(sample_image: Path):
    from PIL import Image

    with Image.open(sample_image) as im:
        original_ratio = im.width / im.height

    data = SlideData(title="T", images=[ImageRef(path=sample_image)])
    prs = build_presentation([data])
    picture = next(s for s in prs.slides[0].shapes if type(s).__name__ == "Picture")
    built_ratio = picture.width / picture.height
    assert abs(built_ratio - original_ratio) < 0.01
